"""Reading an uploaded CSV, JSON or Excel file into header + rows.

Every value comes out as a stripped string ("" for empty): the typed meaning of
a column is only known once it is mapped to a field, so conversion happens in
fields.py, per field, where a bad value can be reported against its row.
"""

import csv
import io
import json
from datetime import date, datetime

from app.core.exceptions import ValidationAppError

MAX_FILE_BYTES = 50 * 1024 * 1024
# Sized so the check step -- one request over every row -- stays well inside
# Cloud Run's 5-minute request timeout on a single CPU. ~100,000 rows checks in
# about 30 seconds locally; larger exports are split into several files.
MAX_ROWS = 150_000

_CSV_EXTENSIONS = {".csv", ".tsv", ".txt"}
_EXCEL_EXTENSIONS = {".xlsx"}
_JSON_EXTENSIONS = {".json"}
SUPPORTED_EXTENSIONS = sorted(_CSV_EXTENSIONS | _EXCEL_EXTENSIONS | _JSON_EXTENSIONS)


def _extension(file_name: str) -> str:
    dot = file_name.rfind(".")
    return file_name[dot:].lower() if dot != -1 else ""


def _clean_headers(raw: list) -> list[str]:
    """Blank headers get a placeholder name and repeated ones a suffix, so every
    column can be told apart when it is mapped."""
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(raw, start=1):
        name = str(value).strip() if value is not None else ""
        name = name or f"Column {index}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        headers.append(name)
    return headers


def _as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, list):
        return ", ".join(_as_text(v) for v in value if v is not None)
    if isinstance(value, dict):
        return json.dumps(value)
    return str(value).strip()


def _decode(data: bytes) -> str:
    # Excel saves "CSV UTF-8" with a byte-order mark and plain "CSV" in the
    # Windows code page; both are common in exports from other systems.
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _read_csv(data: bytes) -> tuple[list[str], list[list[str]]]:
    text = _decode(data)
    sample = text[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        first_line = sample.splitlines()[0] if sample else ""
        delimiter = max(",\t;|", key=first_line.count)
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = list(reader)
    if not rows:
        return [], []
    return _clean_headers(rows[0]), rows[1:]


def _read_excel(data: bytes) -> tuple[list[str], list[list]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationAppError("That file could not be read as an Excel workbook.") from exc
    sheet = workbook.worksheets[0]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    if not rows:
        return [], []
    return _clean_headers(rows[0]), rows[1:]


def _read_json(data: bytes) -> tuple[list[str], list[list]]:
    try:
        parsed = json.loads(_decode(data))
    except json.JSONDecodeError as exc:
        raise ValidationAppError(f"That file is not valid JSON ({exc.msg}, line {exc.lineno}).") from exc
    if isinstance(parsed, dict):
        # Accept the common wrappers: {"rows": [...]}, {"data": [...]}, or a
        # single object whose only value is the list.
        for key in ("rows", "data", "items", "records", "results"):
            if isinstance(parsed.get(key), list):
                parsed = parsed[key]
                break
        else:
            lists = [v for v in parsed.values() if isinstance(v, list)]
            parsed = lists[0] if len(lists) == 1 else [parsed]
    if not isinstance(parsed, list) or not all(isinstance(item, dict) for item in parsed):
        raise ValidationAppError("A JSON import must be a list of objects, one per record.")

    headers: list[str] = []
    for item in parsed:
        for key in item:
            if key not in headers:
                headers.append(key)
    return _clean_headers(headers), [[item.get(h) for h in headers] for item in parsed]


def parse_file(data: bytes, file_name: str) -> tuple[list[str], list[dict[str, str]]]:
    """The file's column headers and its non-empty rows as {header: text}."""
    if not data:
        raise ValidationAppError("The uploaded file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ValidationAppError(
            f"The file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB. Split it and import each part."
        )

    extension = _extension(file_name)
    if extension in _CSV_EXTENSIONS:
        headers, raw_rows = _read_csv(data)
    elif extension in _EXCEL_EXTENSIONS:
        headers, raw_rows = _read_excel(data)
    elif extension in _JSON_EXTENSIONS:
        headers, raw_rows = _read_json(data)
    else:
        raise ValidationAppError(f"Unsupported file type. Use one of: {', '.join(SUPPORTED_EXTENSIONS)}.")

    if not headers:
        raise ValidationAppError("The file has no header row.")

    rows: list[dict[str, str]] = []
    for raw in raw_rows:
        values = [_as_text(v) for v in raw]
        if not any(values):
            continue
        rows.append({header: values[i] if i < len(values) else "" for i, header in enumerate(headers)})
        if len(rows) > MAX_ROWS:
            raise ValidationAppError(f"The file has more than {MAX_ROWS:,} rows. Split it and import each part.")
    if not rows:
        raise ValidationAppError("The file has a header row but no data rows.")
    return headers, rows
